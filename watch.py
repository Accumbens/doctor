#!/usr/bin/env python3
"""
zzzs-watch — daily check for personal doctors (osebni zdravnik) accepting new
patients in the Ljubljana area, straight from the official ZZZS spreadsheet.

Source page:
  https://zavarovanec.zzzs.si/izbira-in-zamenjava-osebnega-zdravnika/seznami-zdravnikov/
File: "Število opredeljenih pri splošnih zdravnikih" (SA_ZO_<date>.xlsx)

Needs: openpyxl  (pip install openpyxl)
"""

import argparse
import csv
import gzip
import smtplib
import html as html_mod
import http.cookiejar
import json
import os
import re
import sys
import time
import unicodedata
import urllib.request
import zlib
from datetime import datetime, timezone
from email.message import EmailMessage

import openpyxl

# --------------------------------------------------------------------------
# CONFIG
# --------------------------------------------------------------------------

LISTING_PAGE = os.environ.get(
    "LISTING_PAGE",
    "https://zavarovanec.zzzs.si/izbira-in-zamenjava-osebnega-zdravnika/"
    "seznami-zdravnikov/")

# Fallback only. The cHash is a TYPO3 hash over the query params and stdatoteke
# is a file ID; both can change when ZZZS republishes. The listing page is
# scraped first precisely so a rotated URL doesn't freeze this silently.
FALLBACK_URL = (
    "https://partner.zzzs.si/zzzs-api/prenesi-datoteko?stdatoteke=14314"
    "&tx_agzzzsapi_izvajalcidownloadfile%5Baction%5D=downloadfile"
    "&tx_agzzzsapi_izvajalcidownloadfile%5Bcontroller%5D=Izvajalci"
    "&type=1249058992&cHash=72bafd008b288be0a7ce04ac87c8f373"
)

# Link label on the listing page, matched case-insensitively on normalised text.
LINK_LABEL_RAW = "Stevilo opredeljenih pri splosnih zdravnikih"

# Ljubljana only: the practice city must contain one of these. Covers the
# 1000 LJUBLJANA postal area plus suburbs like LJUBLJANA - POLJE / - CRNUCE.
# Comma-separate to widen, e.g. CITIES="LJUBLJANA,DOMZALE".
# Normalised lazily in parse(), because norm() is defined further down.
CITIES_RAW = os.environ.get("CITIES", "LJUBLJANA")

# Adult family medicine only. Add "otroski" to include paediatric/school clinics.
WANTED_ACTIVITY = ("splosna ambulanta", "specializanta")

# Column indexes in the sheet (0-based, header on row 4, data from row 5).
C_REGION, C_INST_ID, C_INST, C_STREET, C_CITY = 0, 1, 2, 3, 4
C_DOC_ID, C_DOCTOR, C_ACT_ID, C_ACTIVITY = 5, 6, 7, 8
C_ACCEPTS, C_MUST_INFANTS, C_FTE, C_PATIENTS, C_QUOTIENTS, C_TRAINEE = 9, 10, 11, 12, 13, 14

STATE_FILE = os.environ.get("STATE_FILE", "state.json")
IGNORE_FILE = os.environ.get("IGNORE_FILE", "ignore.txt")

# Optional: a published-to-web Google Sheet CSV. Takes priority over the local
# file so the list can be edited from a phone. Column named "Name" if present,
# otherwise the first column.
IGNORE_URL = os.environ.get("IGNORE_URL", "")

# --------------------------------------------------------------------------
# HELPERS
# --------------------------------------------------------------------------


def norm(s):
    """Uppercase, strip diacritics, collapse whitespace. For robust matching."""
    if not s:
        return ""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().upper()


def name_key(s):
    """Order-independent name key, so 'Dražen Žagar' matches 'ŽAGAR  DRAŽEN'."""
    return " ".join(sorted(norm(s).replace(",", " ").split()))


def _names_from_csv(text):
    rows = list(csv.reader(text.splitlines()))
    if not rows:
        return set()
    header = [norm(c) for c in rows[0]]
    col = header.index("NAME") if "NAME" in header else 0
    out = set()
    for r in rows[1:]:
        if len(r) > col and r[col].strip():
            out.add(name_key(r[col]))
    return out


def load_ignore():
    """Names already tried — from a published Google Sheet, or a local file."""
    if IGNORE_URL:
        try:
            text = http_get(IGNORE_URL, timeout=30)[0].decode("utf-8", "replace")
            names = _names_from_csv(text)
            if names:
                return names
            # An empty result usually means a login page, not an empty sheet.
            print("warn: ignore sheet returned no names; falling back to local file",
                  file=sys.stderr)
        except Exception as e:
            print(f"warn: could not read ignore sheet ({e}); using local file",
                  file=sys.stderr)

    if not os.path.exists(IGNORE_FILE):
        return set()
    out = set()
    with open(IGNORE_FILE, encoding="utf-8") as f:
        for line in f:
            line = line.split("#")[0].strip()
            if line:
                out.add(name_key(line))
    return out


# --------------------------------------------------------------------------
# FETCH
# --------------------------------------------------------------------------


def http_get(url, timeout=90, retries=3):
    """GET with retries, gzip support, and the response headers returned too."""
    last = None
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers={
                "User-Agent": "Mozilla/5.0 (zzzs-watch; personal use)",
                "Accept": "*/*",
                "Accept-Encoding": "gzip, deflate",
                "Accept-Language": "sl,en;q=0.8",
            })
            # Cookie support: TYPO3 sites often redirect once to set a session.
            opener = urllib.request.build_opener(
                urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
            with opener.open(req, timeout=timeout) as r:
                raw, headers = r.read(), r.headers
            enc = (headers.get("Content-Encoding") or "").lower()
            if enc == "gzip":
                raw = gzip.decompress(raw)
            elif enc == "deflate":
                raw = zlib.decompress(raw, -zlib.MAX_WBITS)
            return raw, headers
        except Exception as e:
            last = e
            if attempt < retries - 1:
                time.sleep(2 ** attempt * 3)   # 3s, 6s
    raise last


def resolve_download_url():
    """Find the current xlsx link on the listing page; fall back if not found."""
    try:
        html = http_get(LISTING_PAGE)[0].decode("utf-8", "replace")
    except Exception as e:
        print(f"warn: listing page unreachable ({e}); using fallback URL", file=sys.stderr)
        return FALLBACK_URL

    # Each anchor: capture href and its visible text, then match on the label.
    for href, text in re.findall(r'<a[^>]+href="([^"]+)"[^>]*>(.*?)</a>', html,
                                 re.S | re.I):
        # Strip nested tags, decode entities (&#353; = š), then normalise both
        # sides the same way — norm() uppercases, so the constant must too.
        label = norm(html_mod.unescape(re.sub(r"<[^>]+>", " ", text)))
        if norm(LINK_LABEL_RAW) in label:
            url = html_mod.unescape(href)
            if url.startswith("/"):
                url = "https://zavarovanec.zzzs.si" + url
            return url

    print("warn: link label not found on listing page; using fallback URL",
          file=sys.stderr)
    return FALLBACK_URL


def download(path, verbose=False):
    url = resolve_download_url()
    if verbose:
        print(f"resolved URL: {url}")
    data, headers = http_get(url)

    if verbose:
        print(f"  status content-type: {headers.get('Content-Type')}")
        print(f"  content-disposition: {headers.get('Content-Disposition')}")
        print(f"  bytes: {len(data)}")

    # An xlsx is a zip archive. Anything else means we were served an error or
    # login page, which would otherwise parse as "zero doctors accepting" and
    # silently suppress alerts forever.
    if data[:2] != b"PK":
        snippet = data[:200].decode("utf-8", "replace")
        raise RuntimeError(f"not an xlsx ({len(data)} bytes). Server said: {snippet!r}")

    # ZZZS names the file SA_ZO_<dd_mm_yyyy>.xlsx via Content-Disposition.
    served = re.search(r'filename\*?=(?:UTF-8\'\')?"?([^";]+)',
                       headers.get("Content-Disposition") or "")
    if served and verbose:
        print(f"  served filename: {served.group(1)}")

    with open(path, "wb") as f:
        f.write(data)
    return path


# --------------------------------------------------------------------------
# PARSE
# --------------------------------------------------------------------------


def parse(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=1, values_only=True))

    # "Stanje podatkov na dan:" sits in row 2. Used to spot a frozen feed.
    data_date = ""
    for cell in (rows[1] if len(rows) > 1 else []):
        if isinstance(cell, datetime):
            data_date = cell.strftime("%d.%m.%Y")
        elif isinstance(cell, str) and re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", cell.strip()):
            data_date = cell.strip()

    cities = tuple(norm(c) for c in CITIES_RAW.split(",") if c.strip())

    hits = []
    for r in rows[4:]:
        if not r or r[C_INST_ID] is None:
            continue
        if norm(r[C_ACCEPTS]) != "DA":
            continue

        activity = norm(r[C_ACTIVITY]).lower()
        if not any(w in activity for w in WANTED_ACTIVITY):
            continue

        region = norm(r[C_REGION])
        city = norm(r[C_CITY])
        if not any(c in city for c in cities):
            continue

        hits.append({
            # Doctor code is the stable identity; names get respelled.
            "key": f"{r[C_DOC_ID]}|{r[C_INST_ID]}|{r[C_ACT_ID]}",
            "doctor": re.sub(r"\s+", " ", str(r[C_DOCTOR] or "").strip()),
            "region": region,
            "institution": str(r[C_INST] or "").strip(),
            "address": str(r[C_STREET] or "").strip(),
            "city": str(r[C_CITY] or "").strip(),
            "activity": str(r[C_ACTIVITY] or "").strip(),
            "patients": r[C_PATIENTS],
            "quotients": r[C_QUOTIENTS],
            "fte": r[C_FTE],
            "trainee": str(r[C_TRAINEE] or "").strip(),
            "must_take_infants": norm(r[C_MUST_INFANTS]) == "DA",
        })
    return data_date, hits


# --------------------------------------------------------------------------
# STATE / OUTPUT
# --------------------------------------------------------------------------


def state_is_stale(data_date, max_days=4):
    """True if the published data date is older than max_days. Catches a dead feed."""
    try:
        d = datetime.strptime(data_date, "%d.%m.%Y")
    except ValueError:
        return False
    return (datetime.now() - d).days > max_days


def load_state():
    if not os.path.exists(STATE_FILE):
        return None
    try:
        with open(STATE_FILE, encoding="utf-8") as f:
            return json.load(f)
    except (json.JSONDecodeError, OSError) as e:
        print(f"warn: unreadable {STATE_FILE} ({e}); treating as first run", file=sys.stderr)
        return None


def save_state(keys, data_date):
    tmp = STATE_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump({"checked_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                   "data_date": data_date,
                   "accepting": sorted(keys)}, f, ensure_ascii=False, indent=1)
    os.replace(tmp, STATE_FILE)


def fmt(h):
    lines = [f"• {h['doctor']}"]
    lines.append(f"  {h['institution']}")
    where = ", ".join(x for x in (h["address"], h["city"]) if x)
    if where:
        lines.append(f"  {where}")
    if h["patients"] is not None:
        # Below ~1895 quotients the doctor is still obliged to accept.
        lines.append(f"  {h['patients']:.0f} patients / {h['quotients']:.0f} quotients "
                     f"(FTE {h['fte']})")
    if h["trainee"]:
        lines.append(f"  specializant: {h['trainee']}")
    return "\n".join(lines)


def notify_ntfy(title, body):
    topic = os.environ.get("NTFY_TOPIC")
    if not topic:
        return False
    server = os.environ.get("NTFY_SERVER", "https://ntfy.sh").rstrip("/")
    req = urllib.request.Request(f"{server}/{topic}", data=body.encode(),
                                 headers={"Priority": "high"})
    # ntfy headers are latin-1 only; Slovene names would raise otherwise.
    req.add_header("Title", title.encode("utf-8").decode("latin-1", "replace"))
    if os.environ.get("NTFY_TOKEN"):
        req.add_header("Authorization", f"Bearer {os.environ['NTFY_TOKEN']}")
    with urllib.request.urlopen(req, timeout=30):
        pass
    return True


def notify_email(title, body):
    to_addr = os.environ.get("EMAIL_TO")
    host = os.environ.get("SMTP_HOST", "smtp.gmail.com")
    if not to_addr:
        return False

    msg = EmailMessage()
    msg["Subject"] = title            # EmailMessage handles UTF-8 headers itself
    msg["From"] = os.environ.get("EMAIL_FROM") or os.environ.get("SMTP_USER") or to_addr
    msg["To"] = to_addr
    msg.set_content(body)

    port = int(os.environ.get("SMTP_PORT", "587"))
    user, pw = os.environ.get("SMTP_USER"), os.environ.get("SMTP_PASS")

    if port == 465:
        smtp = smtplib.SMTP_SSL(host, port, timeout=30)
    else:
        smtp = smtplib.SMTP(host, port, timeout=30)
    with smtp as s:
        if port != 465 and os.environ.get("SMTP_STARTTLS", "1") == "1":
            s.starttls()
        if user and pw:
            s.login(user, pw)
        s.send_message(msg)
    return True


def notify_file(title, body):
    """Write the alert to disk so CI can turn it into a GitHub issue (which
    GitHub then emails you). Needs no credentials at all."""
    path = os.environ.get("ALERT_FILE")
    if not path:
        return False
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"{title}\n\n{body}\n")
    return True


def notify(title, body):
    """Fire every configured channel; one failing must not silence the others."""
    sent = 0
    for fn in (notify_ntfy, notify_email, notify_file):
        try:
            if fn(title, body):
                sent += 1
        except Exception as e:
            print(f"warn: {fn.__name__} failed: {e}", file=sys.stderr)
    if sent == 0:
        print(f"\n=== {title} ===\n{body}\n")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--file", help="parse a local xlsx instead of downloading")
    ap.add_argument("--list", action="store_true", help="print current matches, keep state")
    ap.add_argument("--all", action="store_true", help="don't apply the ignore list")
    ap.add_argument("--alert-on-first-run", action="store_true")
    ap.add_argument("--test-notify", action="store_true",
                    help="send a test message through every configured channel")
    ap.add_argument("--check-fetch", action="store_true",
                    help="download and report what the server returned, then exit")
    args = ap.parse_args()

    if args.test_notify:
        notify("zzzs-watch test",
               "If you can read this, notifications are working.")
        return 0

    if args.check_fetch:
        p = download("SA_ZO_latest.xlsx", verbose=True)
        d, h = parse(p)
        print(f"parsed OK: data dated {d}, {len(h)} matching doctors")
        return 0

    # Everything below is the daily run. It always sends exactly one
    # notification — new doctors, "nothing new", or "something broke" — so a
    # missing email is never ambiguous between "quiet day" and "silently dead".
    try:
        path = args.file or download("SA_ZO_latest.xlsx")
        data_date, hits = parse(path)

        stale_note = ""
        if data_date and state_is_stale(data_date):
            stale_note = (f"\n\nWarning: ZZZS data is still dated {data_date} "
                          "— the feed may be frozen upstream.")
            print(f"warn: ZZZS data still dated {data_date} — feed may be frozen",
                  file=sys.stderr)

        ignored = set() if args.all else load_ignore()
        skipped = [h for h in hits if name_key(h["doctor"]) in ignored]
        visible = [h for h in hits if name_key(h["doctor"]) not in ignored]
        visible.sort(key=lambda h: (h["patients"] is None, h["patients"]))

        if args.list:
            print(f"ZZZS data dated {data_date} — {len(visible)} accepting"
                  f"{f', {len(skipped)} ignored' if skipped else ''}\n")
            for h in visible:
                print(fmt(h), "\n")
            return 0

        state = load_state()
        previous = set(state["accepting"]) if state else set()
        first_run = state is None
        # Diff against EVERY currently-accepting doctor, ignored or not, so
        # state remembers who's already been seen regardless of ignore status.
        # Filtering by "ignored" first would mean un-ignoring someone (or a
        # transient ignore-list fetch failure) makes them look "new" again
        # even though they've been accepting the whole time.
        new = [h for h in hits if h["key"] not in previous
               and name_key(h["doctor"]) not in ignored]

        print(f"[{datetime.now():%Y-%m-%d %H:%M}] data {data_date} | matching {len(visible)} "
              f"| new {len(new)} | ignored {len(skipped)}"
              + ("  (first run)" if first_run else ""))

        if new and (not first_run or args.alert_on_first_run):
            title = (f"{len(new)} new doctor(s) accepting near Ljubljana"
                     if len(new) > 1 else f"{new[0]['doctor']} is accepting")
            body = "\n\n".join(fmt(h) for h in new)
            body += f"\n\nZZZS data dated {data_date}. Phone before you travel."
        elif first_run:
            title = "zzzs-watch: first run — baseline recorded"
            body = (f"Recorded {len(visible)} doctor(s) currently accepting as the "
                    f"starting point ({len(skipped)} more matched but are already on "
                    "your ignore list, so weren't counted above). "
                    "You'll hear about anything new from the next run.\n\n"
                    f"ZZZS data dated {data_date}.")
        else:
            title = "zzzs-watch: no new doctors today"
            body = (f"{len(visible)} doctor(s) accepting, none new since yesterday "
                    f"({len(skipped)} more matched but are on your ignore list, so "
                    "aren't counted above).\n\n"
                    f"ZZZS data dated {data_date}.")
        notify(title, body + stale_note)

        save_state([h["key"] for h in hits], data_date)
        return 0

    except Exception as e:
        print(f"error: {type(e).__name__}: {e}", file=sys.stderr)
        notify("zzzs-watch: problem running today's check",
               f"Today's check failed before it could finish, so treat this as "
               f"'unknown' rather than 'nothing new':\n\n{type(e).__name__}: {e}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
